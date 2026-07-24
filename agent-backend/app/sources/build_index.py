"""Build SQLite indexes for large local tables (gene / site keyed lookup).

Usage:
  python -m app.sources.build_index activedriverdb
  python -m app.sources.build_index activedriverdb --file clinvar
  python -m app.sources.build_index --all
"""

from __future__ import annotations

import argparse
import csv
import logging
import re
import sqlite3
from pathlib import Path
from typing import Any, Iterable

from app.sources.catalog import get_catalog, reload_catalog
from app.sources.models import DataFile, SourceManifest

logger = logging.getLogger(__name__)

# Header aliases: raw TSV/CSV headers → normalized column names
_HEADER_ALIASES: dict[str, str] = {
    "gene": "gene",
    "genename": "gene",
    "refseq": "refseq",
    "mutation position": "mutation_position",
    "mutation alt": "mutation_alt",
    "mutation summary": "mutation_summary",
    "site position": "site_position",
    "site residue": "site_residue",
    "kinase symbol": "kinase_symbol",
    "target symbol": "target_symbol",
    "kinase refseq": "kinase_refseq",
    "target refseq": "target_refseq",
    "target sequence position": "target_sequence_position",
    "target amino acid": "target_amino_acid",
    # CancerProteome
    "up": "uniprot",
    "pos": "position",
    "mods": "ptm_ome",
    "mean expression in control samples": "mean_control",
    "mean expression in tumor samples": "mean_tumor",
    # PhosphoSitePlus
    "acc_id": "uniprot",
    "sub_acc_id": "sub_uniprot",
    "kin_acc_id": "kin_uniprot",
    "disease(s)": "diseases",
    "mut_rsd#": "mut_position",
    "aa_change": "aa_change",
    "var_type": "var_type",
    "mod_type": "mod_type",
    "site_+/-7_aa": "site_flank",
    "pmids": "pmids",
    # dSCOPE
    "uniid": "uniprot",
    "uniprot_id": "uniprot",
    "averagescores": "averagescores",
    "regions": "regions",
}


_MOD_RSD_RE = re.compile(r"^([A-Za-z])?(\d+)(?:-([A-Za-z0-9]+))?$")


def normalize_header(name: str) -> str:
    key = name.strip().lower()
    if key in _HEADER_ALIASES:
        return _HEADER_ALIASES[key]
    return key.replace(" ", "_").replace("#", "").replace("(", "").replace(")", "")


def index_path_for(manifest: SourceManifest, file_meta: DataFile) -> Path:
    assert manifest.root is not None
    return manifest.root / "indexes" / f"{file_meta.id}.sqlite"


def _parse_mod_position(mod_rsd: str) -> str:
    """Extract residue position from S15-p, S15, or bare 15."""
    text = (mod_rsd or "").strip()
    if not text:
        return ""
    m = _MOD_RSD_RE.match(text)
    return m.group(2) if m else ""


def _delimiter(file_meta: DataFile) -> str:
    delim = file_meta.delimiter or "\t"
    if delim == "\\t":
        return "\t"
    return delim


def _is_text_tabular(src: Path) -> bool:
    """Detect tab-/comma-separated text files with spreadsheet-like extensions."""
    with src.open("rb") as f:
        head = f.read(512)
    if head.startswith(b"PK"):  # xlsx / zip container
        return False
    try:
        text = head.decode("utf-8", errors="replace")
    except Exception:
        return False
    return "\t" in text or "," in text


def _iter_text_rows(src: Path, file_meta: DataFile) -> tuple[list[str], Iterable[list[str]]]:
    """Return (raw_header_cells, data_row_iterator) for text/CSV/TSV files."""
    encoding = file_meta.encoding or "utf-8"
    delim = _delimiter(file_meta)
    f = src.open(encoding=encoding, newline="", errors="replace")
    reader = csv.reader(f, delimiter=delim)

    raw_header: list[str] | None = None
    if file_meta.skip_until_header:
        marker = (file_meta.header_startswith or "").strip()
        for row in reader:
            if not row or all(not (c or "").strip() for c in row):
                continue
            first = (row[0] or "").strip()
            # Skip PSP date / license prose
            if first.startswith("#"):
                continue
            if marker:
                if first == marker or (row and delim.join(row).startswith(marker)):
                    raw_header = row
                    break
            elif "\t" in delim.join(row) or len(row) > 3:
                # Heuristic: first wide tabular row is the header
                if first.upper() in {
                    "GENE", "DISEASE", "PROTEIN", "KINASE", "ORGANISM", "UPID",
                } or first.upper().endswith("_ID"):
                    raw_header = row
                    break
        if raw_header is None:
            f.close()
            raise ValueError(f"Could not find header row in {src}")
    else:
        raw_header = next(reader, None)
        if not raw_header:
            f.close()
            raise ValueError(f"Empty header in {src}")

    def _rows() -> Iterable[list[str]]:
        try:
            for row in reader:
                yield row
        finally:
            f.close()

    return raw_header, _rows()


def _iter_xlsx_rows(src: Path, file_meta: DataFile) -> tuple[list[str], Iterable[list[str]]]:
    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError("openpyxl required to index .xlsx files") from e

    wb = openpyxl.load_workbook(str(src), read_only=True, data_only=True)
    sheet_name = file_meta.xlsx_sheet or wb.sheetnames[0]
    ws = wb[sheet_name]
    marker = (file_meta.header_startswith or "GENE").strip().upper()

    raw_header: list[str] | None = None
    pending: list[list[str]] = []
    it = ws.iter_rows(values_only=True)
    for row in it:
        cells = [("" if c is None else str(c)) for c in row]
        if not any(c.strip() for c in cells):
            continue
        first = cells[0].strip().upper()
        if file_meta.skip_until_header:
            if first == marker:
                raw_header = cells
                break
            continue
        raw_header = cells
        break

    if raw_header is None:
        wb.close()
        raise ValueError(f"Could not find header row in {src} sheet={sheet_name}")

    def _rows() -> Iterable[list[str]]:
        try:
            for row in it:
                yield [("" if c is None else str(c)) for c in row]
        finally:
            wb.close()

    return raw_header, _rows()


def build_file_index(manifest: SourceManifest, file_meta: DataFile) -> Path:
    """Stream a table into a SQLite file with indexes on declared keys."""
    assert manifest.root is not None
    src = manifest.resolve(file_meta.path)
    if not src.exists():
        raise FileNotFoundError(src)

    out = index_path_for(manifest, file_meta)
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        out.unlink()

    if src.suffix.lower() in {".xlsx", ".xlsm", ".xls"} and not _is_text_tabular(src):
        raw_header, row_iter = _iter_xlsx_rows(src, file_meta)
    else:
        raw_header, row_iter = _iter_text_rows(src, file_meta)

    cols = [normalize_header(h) for h in raw_header]
    # Drop empty trailing header names
    while cols and not cols[-1]:
        cols.pop()
        raw_header = raw_header[: len(cols)]
    if not cols:
        raise ValueError(f"Empty header in {src}")

    # Deduplicate column names if needed
    seen: dict[str, int] = {}
    uniq_cols: list[str] = []
    for c in cols:
        if c not in seen:
            seen[c] = 0
            uniq_cols.append(c)
        else:
            seen[c] += 1
            uniq_cols.append(f"{c}_{seen[c]}")
    cols = uniq_cols

    derive_uniprot = (
        "protein" in cols
        and "uniprot" not in cols
        and file_meta.id == "protein"
    )
    derive_pos_col = file_meta.derive_mod_position_from
    if derive_pos_col:
        derive_pos_col = normalize_header(derive_pos_col)

    extra: list[str] = []
    if derive_uniprot:
        extra.append("uniprot")
    if derive_pos_col and "position" not in cols:
        extra.append("position")

    all_cols = list(cols) + extra
    col_sql = ", ".join(f'"{c}" TEXT' for c in all_cols)
    placeholders = ", ".join("?" for _ in all_cols)

    conn = sqlite3.connect(str(out))
    try:
        conn.execute(f"CREATE TABLE records ({col_sql})")
        batch: list[tuple[Any, ...]] = []
        n = 0
        protein_idx = cols.index("protein") if "protein" in cols else -1
        mod_idx = cols.index(derive_pos_col) if derive_pos_col and derive_pos_col in cols else -1
        base_n = len(cols)

        for row in row_iter:
            if not row or all(not (c or "").strip() for c in row):
                continue
            if len(row) < base_n:
                row = list(row) + [""] * (base_n - len(row))
            elif len(row) > base_n:
                row = list(row[:base_n])
            else:
                row = list(row)

            extras_vals: list[str] = []
            if derive_uniprot:
                prot = row[protein_idx] if protein_idx >= 0 else ""
                extras_vals.append(prot.rsplit("_", 1)[-1] if "_" in prot else "")
            if derive_pos_col and "position" in extra:
                mod = row[mod_idx] if mod_idx >= 0 else ""
                extras_vals.append(_parse_mod_position(mod))

            batch.append(tuple(row + extras_vals))
            n += 1
            if len(batch) >= 5000:
                conn.executemany(
                    f"INSERT INTO records VALUES ({placeholders})", batch
                )
                batch.clear()
        if batch:
            conn.executemany(
                f"INSERT INTO records VALUES ({placeholders})", batch
            )

        for key in file_meta.index_keys:
            key_n = normalize_header(key)
            if key_n in all_cols:
                conn.execute(
                    f'CREATE INDEX IF NOT EXISTS idx_{key_n} ON records("{key_n}")'
                )
        conn.commit()
        logger.info(
            "Indexed %s → %s (%d rows, keys=%s)",
            src.name, out.name, n, file_meta.index_keys,
        )
    finally:
        conn.close()
    return out


def build_source_indexes(source_id: str, file_id: str | None = None) -> list[Path]:
    catalog = reload_catalog()
    manifest = catalog.require(source_id)
    built: list[Path] = []
    for meta in manifest.files:
        if file_id and meta.id != file_id:
            continue
        if not meta.index_keys:
            logger.info("Skip %s (no index_keys)", meta.id)
            continue
        built.append(build_file_index(manifest, meta))
    return built


def main(argv: list[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    parser = argparse.ArgumentParser(description="Build local SQLite indexes for PTM sources")
    parser.add_argument("source_id", nargs="?", help="Source id from SOURCE.yaml")
    parser.add_argument("--file", dest="file_id", help="Only one file id")
    parser.add_argument("--all", action="store_true", help="Index every source with files")
    args = parser.parse_args(argv)

    catalog = get_catalog()
    if args.all:
        ids = [s.id for s in catalog.all() if s.files]
    elif args.source_id:
        ids = [args.source_id]
    else:
        parser.error("Provide source_id or --all")
        return 2

    for sid in ids:
        paths = build_source_indexes(sid, args.file_id)
        for p in paths:
            print(p)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
